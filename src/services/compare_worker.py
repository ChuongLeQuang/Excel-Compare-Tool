import logging
import openpyxl
import pandas as pd
from PyQt6.QtCore import QThread, pyqtSignal

from src.models.compare_config import CompareConfig
from src.services.compare_service import ExcelCompareService

logger = logging.getLogger(__name__)


class CompareWorker(QThread):
    """
    EN: Background worker thread to prevent GUI freezing during heavy Excel processing.
    VI: Luồng xử lý nền giúp giao diện không bị đơ khi xử lý file Excel nặng.
    """
    progress = pyqtSignal(int, str)
    finished_success = pyqtSignal(object, object)  # df_src_res, df_cmp_res
    finished_error = pyqtSignal(str)

    def __init__(self, config: CompareConfig, src_path: str, src_sheet: str, src_header: int, cmp_path: str, cmp_sheet: str, cmp_header: int):
        super().__init__()
        self.config = config
        self.src_path = src_path
        self.src_sheet = src_sheet
        self.src_header = src_header
        self.cmp_path = cmp_path
        self.cmp_sheet = cmp_sheet
        self.cmp_header = cmp_header

    def _update_excel_in_place(self, file_path: str, sheet_name: str, df: pd.DataFrame, header_row: int, update_cols: set):
        """
        EN: Update targeted columns directly into the original Excel file preserving formats.
        VI: Ghi trực tiếp các cột kết quả vào file Excel gốc, bảo toàn định dạng cũ.
        """
        wb = openpyxl.load_workbook(file_path)
        sheet = wb[sheet_name]
        
        col_name_to_idx = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=header_row, column=col_idx).value
            if val is not None:
                col_name_to_idx[str(val)] = col_idx

        for col_name in update_cols:
            if col_name not in col_name_to_idx:
                new_col_idx = (sheet.max_column + 1) if sheet.max_column else 1
                sheet.cell(row=header_row, column=new_col_idx).value = col_name
                col_name_to_idx[col_name] = new_col_idx

        for row_idx in range(len(df)):
            excel_row = header_row + 1 + row_idx
            for col_name in update_cols:
                if col_name in df.columns:
                    val = df.at[row_idx, col_name]
                    sheet.cell(row=excel_row, column=col_name_to_idx[col_name]).value = "" if pd.isna(val) else val
        wb.save(file_path)
        wb.close()

    def run(self):
        try:
            logger.info("Worker thread started processing.")
            self.progress.emit(10, "Đang khởi tạo cấu hình và đọc file Excel...")
            service = ExcelCompareService(self.config)
            
            self.progress.emit(50, "Đang xử lý logic đối chiếu dữ liệu...")
            df_src_res, df_cmp_res = service.process_comparison()
            
            src_update_cols = set(r.target_col for r in self.config.rules.on_match_source + self.config.rules.on_unmatch_source)
            cmp_update_cols = set(r.target_col for r in self.config.rules.on_match_compare + self.config.rules.on_unmatch_compare)

            if src_update_cols:
                self.progress.emit(70, f"Đang ghi đè kết quả vào file Source...")
                self._update_excel_in_place(self.src_path, self.src_sheet, df_src_res, self.src_header, src_update_cols)
            if cmp_update_cols:
                self.progress.emit(85, f"Đang ghi đè kết quả vào file Compare...")
                self._update_excel_in_place(self.cmp_path, self.cmp_sheet, df_cmp_res, self.cmp_header, cmp_update_cols)

            self.progress.emit(100, "Hoàn tất!")
            logger.info("Worker thread finished successfully.")
            self.finished_success.emit(df_src_res, df_cmp_res)
        except PermissionError:
            logger.error("PermissionError: Excel file is currently open.")
            self.finished_error.emit("PERMISSION_ERROR")
        except Exception as e:
            logger.error(f"Worker thread error: {str(e)}")
            self.finished_error.emit(str(e))